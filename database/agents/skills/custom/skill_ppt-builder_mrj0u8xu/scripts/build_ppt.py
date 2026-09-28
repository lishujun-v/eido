import argparse
import json
import os
from dataclasses import dataclass
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class Theme:
    accent_rgb: RGBColor
    font_family: str


def _hex_to_rgb(hex_color: str) -> RGBColor:
    v = hex_color.strip().lstrip("#")
    if len(v) != 6:
        return RGBColor(0x2F, 0x6F, 0xED)
    r = int(v[0:2], 16)
    g = int(v[2:4], 16)
    b = int(v[4:6], 16)
    return RGBColor(r, g, b)


def _get_theme(spec: Dict[str, Any]) -> Theme:
    theme = (spec.get("meta") or {}).get("theme") or {}
    accent = _hex_to_rgb(str(theme.get("accentColor") or "#2F6FED"))
    font = str(theme.get("fontFamily") or "Calibri")
    return Theme(accent_rgb=accent, font_family=font)


def _add_title(slide, title: str, theme: Theme) -> None:
    tf = slide.shapes.title.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = theme.accent_rgb
    run.font.name = theme.font_family


def _add_subtitle(slide, subtitle: str, theme: Theme) -> None:
    if len(slide.placeholders) < 2:
        return
    ph = slide.placeholders[1]
    tf = ph.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = subtitle
    run.font.size = Pt(20)
    run.font.name = theme.font_family


def _set_body_bullets(shape, bullets: List[str], theme: Theme) -> None:
    tf = shape.text_frame
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(b)
        p.level = 0
        p.font.size = Pt(22)
        p.font.name = theme.font_family


def _safe_join(root: str, rel_path: str) -> str:
    p = os.path.normpath(os.path.join(root, rel_path))
    if os.path.commonpath([os.path.abspath(root), os.path.abspath(p)]) != os.path.abspath(root):
        raise ValueError(f"Illegal path outside assetsRoot: {rel_path}")
    return p


def _add_picture(slide, image_path: str, left: float, top: float, width: float, height: float) -> None:
    slide.shapes.add_picture(image_path, Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def build_ppt(spec: Dict[str, Any], out_path: str) -> None:
    prs = Presentation()

    theme = _get_theme(spec)
    assets_root = str(spec.get("assetsRoot") or "assets")

    slides = spec.get("slides") or []
    if not isinstance(slides, list):
        raise ValueError("spec.slides must be a list")

    for s in slides:
        layout = str((s or {}).get("layout") or "bullets")
        title = str((s or {}).get("title") or "")

        if layout == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            _add_title(slide, title or str((spec.get("meta") or {}).get("title") or "演示文稿"), theme)
            subtitle = str((s or {}).get("subtitle") or (spec.get("meta") or {}).get("subtitle") or "")
            if subtitle:
                _add_subtitle(slide, subtitle, theme)
            continue

        if layout == "section":
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            if slide.shapes.title:
                slide.shapes.title.text = title
            continue

        if layout == "quote":
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            if slide.shapes.title:
                slide.shapes.title.text = title or "关键结论"
            quote = str((s or {}).get("quote") or "")
            box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.0))
            tf = box.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = quote
            run.font.size = Pt(34)
            run.font.bold = True
            run.font.name = theme.font_family
            attribution = str((s or {}).get("attribution") or "")
            if attribution:
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                run2 = p2.add_run()
                run2.text = attribution
                run2.font.size = Pt(18)
                run2.font.name = theme.font_family
            continue

        if layout == "picture":
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            if slide.shapes.title:
                slide.shapes.title.text = title
            img = (s or {}).get("image") or {}
            rel = str(img.get("path") or "")
            if rel:
                img_path = _safe_join(assets_root, rel)
                if os.path.exists(img_path):
                    _add_picture(slide, img_path, left=0.8, top=1.6, width=7.2, height=5.2)
            bullets = (s or {}).get("bullets") or []
            if bullets:
                box = slide.shapes.add_textbox(Inches(8.3), Inches(1.6), Inches(4.0), Inches(5.2))
                _set_body_bullets(box, [str(x) for x in bullets], theme)
            caption = str(img.get("caption") or "")
            if caption:
                cap = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(7.2), Inches(0.5))
                tf = cap.text_frame
                tf.text = caption
                tf.paragraphs[0].font.size = Pt(14)
                tf.paragraphs[0].font.name = theme.font_family
            continue

        if layout == "two-column":
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            if slide.shapes.title:
                slide.shapes.title.text = title
            left = (s or {}).get("left") or {}
            right = (s or {}).get("right") or {}

            lbox = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.6))
            ltf = lbox.text_frame
            ltf.clear()
            lhead = str(left.get("heading") or "")
            if lhead:
                p = ltf.paragraphs[0]
                p.text = lhead
                p.font.bold = True
                p.font.size = Pt(24)
                p.font.name = theme.font_family
            lbs = [str(x) for x in (left.get("bullets") or [])]
            if lbs:
                if lhead:
                    ltf.add_paragraph()
                for i, b in enumerate(lbs):
                    p = ltf.add_paragraph() if (lhead or i > 0) else ltf.paragraphs[0]
                    p.text = b
                    p.level = 0
                    p.font.size = Pt(20)
                    p.font.name = theme.font_family

            rbox = slide.shapes.add_textbox(Inches(6.7), Inches(1.6), Inches(5.8), Inches(5.6))
            rtf = rbox.text_frame
            rtf.clear()
            rhead = str(right.get("heading") or "")
            if rhead:
                p = rtf.paragraphs[0]
                p.text = rhead
                p.font.bold = True
                p.font.size = Pt(24)
                p.font.name = theme.font_family
            rbs = [str(x) for x in (right.get("bullets") or [])]
            if rbs:
                if rhead:
                    rtf.add_paragraph()
                for i, b in enumerate(rbs):
                    p = rtf.add_paragraph() if (rhead or i > 0) else rtf.paragraphs[0]
                    p.text = b
                    p.level = 0
                    p.font.size = Pt(20)
                    p.font.name = theme.font_family
            continue

        # default: bullets
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title:
            slide.shapes.title.text = title
        body = slide.shapes.placeholders[1]
        bullets = [str(x) for x in ((s or {}).get("bullets") or [])]
        _set_body_bullets(body, bullets, theme)
        footnote = str((s or {}).get("footnote") or "")
        if footnote:
            fn = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(12.0), Inches(0.4))
            t = fn.text_frame
            t.text = footnote
            t.paragraphs[0].font.size = Pt(12)
            t.paragraphs[0].font.name = theme.font_family

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    prs.save(out_path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec", required=True, help="Path to deck.json")
    ap.add_argument("--out", required=True, help="Output .pptx path")
    args = ap.parse_args()

    with open(args.spec, "r", encoding="utf-8") as f:
        spec = json.load(f)

    build_ppt(spec, args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

